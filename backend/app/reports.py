from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO
from sqlalchemy.orm import Session
from sqlalchemy import func
from app import models
from datetime import datetime
from typing import Optional

def reshape_arabic(text: str) -> str:
    if not text: return ""
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        reshaped = arabic_reshaper.reshape(str(text))
        return get_display(reshaped)
    except ImportError:
        return str(text)

# ==========================================
# Summary Report (PDF & Excel)
# ==========================================
def generate_summary_pdf(db: Session, from_date: datetime = None, to_date: datetime = None) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], alignment=1, fontSize=20, textColor=colors.HexColor('#1F4E78'), spaceAfter=20, fontName='Helvetica-Bold')
    subtitle_style = ParagraphStyle('CustomSubtitle', parent=styles['Normal'], alignment=1, fontSize=12, textColor=colors.HexColor('#7f8c8d'), spaceAfter=30)
    section_style = ParagraphStyle('SectionTitle', parent=styles['Heading2'], fontSize=14, textColor=colors.HexColor('#2c3e50'), spaceAfter=10, spaceBefore=15, fontName='Helvetica-Bold')

    elements.append(Paragraph("NRC Latrine Tracker", title_style))
    elements.append(Paragraph("Project Executive Summary & Governance Report", subtitle_style))
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", subtitle_style))
    elements.append(Spacer(1, 20))

    total = db.query(models.Latrine).count()
    completed = db.query(models.Latrine).filter(models.Latrine.status == 'completed').count()
    in_progress = db.query(models.Latrine).filter(models.Latrine.status == 'in_progress').count()
    not_started = db.query(models.Latrine).filter(models.Latrine.status == 'not_started').count()

    decisions = db.query(models.DecisionRecord).all()
    approved_count = sum(1 for d in decisions if (d.human_decision_code or d.system_recommendation_code) in ['APPROVE', 'APPROVE_WITH_NOTE'])
    hold_count = sum(1 for d in decisions if (d.human_decision_code or d.system_recommendation_code) == 'HOLD')
    rework_count = sum(1 for d in decisions if (d.human_decision_code or d.system_recommendation_code) == 'REWORK')
    stop_count = sum(1 for d in decisions if (d.human_decision_code or d.system_recommendation_code) == 'STOP')

    avg_progress = round(float(db.query(func.avg(models.Latrine.overall_pct)).scalar() or 0), 2)

    elements.append(Paragraph("1. Project Status", section_style))
    summary_data = [
        ['Indicator', 'Value', 'Percentage'],
        ['Total Latrines', str(total), '100%'],
        ['Completed', str(completed), f'{round(completed/total*100,1)}%' if total else '0%'],
        ['In Progress', str(in_progress), f'{round(in_progress/total*100,1)}%' if total else '0%'],
        ['Not Started', str(not_started), f'{round(not_started/total*100,1)}%' if total else '0%'],
        ['Overall Financial Progress', f'{avg_progress}%', 'Based on Approved Payments'],
    ]
    t1 = Table(summary_data, colWidths=[7*cm, 4*cm, 5*cm])
    t1.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
    ]))
    elements.append(t1)
    elements.append(Spacer(1, 20))

    elements.append(Paragraph("2. Decision & Governance Engine", section_style))
    gov_data = [
        ['Decision Status', 'Item Count', 'Action Required'],
        ['Approved for Payment', str(approved_count), 'None'],
        ['On Hold (Pending/Minor Issues)', str(hold_count), 'Review & Inspect'],
        ['Rework Required (Failed Quality)', str(rework_count), 'Contractor Action'],
        ['Stopped (Critical Safety/Quality)', str(stop_count), 'Urgent PM Intervention'],
    ]
    t2 = Table(gov_data, colWidths=[7*cm, 3*cm, 6*cm])
    t2.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#8e44ad')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
    ]))
    elements.append(t2)

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()

def generate_summary_excel(db: Session, from_date: datetime = None, to_date: datetime = None) -> bytes:
    """نفس محتوى الملخص التنفيذي لكن بصيغة Excel"""
    buffer = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    ws.sheet_view.rightToLeft = True

    total = db.query(models.Latrine).count()
    completed = db.query(models.Latrine).filter(models.Latrine.status == 'completed').count()
    in_progress = db.query(models.Latrine).filter(models.Latrine.status == 'in_progress').count()
    not_started = db.query(models.Latrine).filter(models.Latrine.status == 'not_started').count()

    decisions = db.query(models.DecisionRecord).all()
    approved_count = sum(1 for d in decisions if (d.human_decision_code or d.system_recommendation_code) in ['APPROVE', 'APPROVE_WITH_NOTE'])
    hold_count = sum(1 for d in decisions if (d.human_decision_code or d.system_recommendation_code) == 'HOLD')
    rework_count = sum(1 for d in decisions if (d.human_decision_code or d.system_recommendation_code) == 'REWORK')
    stop_count = sum(1 for d in decisions if (d.human_decision_code or d.system_recommendation_code) == 'STOP')

    avg_progress = round(float(db.query(func.avg(models.Latrine.overall_pct)).scalar() or 0), 2)

    # العنوان
    ws.merge_cells('A1:C1')
    ws['A1'] = "NRC Latrine Tracker - Executive Summary"
    ws['A1'].font = Font(size=16, bold=True, color="1F4E78")
    ws['A1'].alignment = Alignment(horizontal='center')

    # جدول حالة المشروع
    ws.cell(row=3, column=1, value="المؤشر").font = Font(bold=True)
    ws.cell(row=3, column=2, value="القيمة").font = Font(bold=True)
    ws.cell(row=3, column=3, value="النسبة").font = Font(bold=True)

    data_rows = [
        ['Total Latrines', total, '100%'],
        ['Completed', completed, f'{round(completed/total*100,1)}%' if total else '0%'],
        ['In Progress', in_progress, f'{round(in_progress/total*100,1)}%' if total else '0%'],
        ['Not Started', not_started, f'{round(not_started/total*100,1)}%' if total else '0%'],
        ['Overall Financial Progress', f'{avg_progress}%', ''],
    ]
    for i, row_data in enumerate(data_rows, start=4):
        for j, val in enumerate(row_data, start=1):
            ws.cell(row=i, column=j, value=val)

    # جدول الحوكمة
    ws.cell(row=10, column=1, value="حالة القرار").font = Font(bold=True)
    ws.cell(row=10, column=2, value="العدد").font = Font(bold=True)
    ws.cell(row=10, column=3, value="الإجراء").font = Font(bold=True)

    gov_data = [
        ['Approved for Payment', approved_count, 'None'],
        ['On Hold', hold_count, 'Review & Inspect'],
        ['Rework Required', rework_count, 'Contractor Action'],
        ['Stopped', stop_count, 'Urgent PM Intervention'],
    ]
    for i, row_data in enumerate(gov_data, start=11):
        for j, val in enumerate(row_data, start=1):
            ws.cell(row=i, column=j, value=val)

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 30

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# ==========================================
# IPC Report (Excel & PDF)
# ==========================================
def generate_ipc_excel(db: Session) -> bytes:
    """شهادة دفع مؤقتة Excel (IPC) - تجميعي"""
    buffer = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "IPC - Master Aggregation"
    ws.sheet_view.rightToLeft = True

    ws['A1'] = "NRC Latrine Tracker - Governance IPC"
    ws['A1'].font = Font(size=16, bold=True, color="1F4E78")
    ws.merge_cells('A1:I1')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    ws['A2'] = "Interim Payment Certificate (Aggregated by BoQ)"
    ws['A2'].font = Font(size=12, bold=True, color="27ae60")
    ws.merge_cells('A2:I2')
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')

    ws['A3'] = f"Date: {datetime.now().strftime('%Y-%m-%d')}"
    ws.merge_cells('A3:I3')
    ws['A3'].alignment = Alignment(horizontal='center')

    headers = [
        'BoQ Code', 'Description (AR)', 'Unit', 'Unit Price ($)',
        'Total Planned Qty', 'Total Executed Qty',
        'Approved Qty (Passed)', 'Payable Amount ($)', 'Blocked Amount ($)'
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    ws.row_dimensions[5].height = 35

    dictionary = {d.boq_code: d for d in db.query(models.BoqDictionary).filter(models.BoqDictionary.is_active == True).all()}
    items = db.query(models.BoqItem).all()
    decisions = {d.boq_item_id: d for d in db.query(models.DecisionRecord).all()}

    aggregation = {}
    for item in items:
        code = item.boq_code
        if code not in aggregation:
            aggregation[code] = {'planned': 0.0, 'executed': 0.0, 'approved_qty': 0.0, 'payable_usd': 0.0, 'blocked_usd': 0.0}

        dict_item = dictionary.get(code)
        price = dict_item.unit_price if dict_item else 0.0
        planned = item.planned_qty or 0.0
        executed = item.achieved_qty or 0.0

        aggregation[code]['planned'] += planned
        aggregation[code]['executed'] += executed

        decision = decisions.get(item.id)
        payment_pct = 0.0
        if decision:
            payment_pct = decision.human_payment_pct if decision.human_payment_pct is not None else decision.system_payment_pct

        payable_qty = planned * (payment_pct / 100.0)
        aggregation[code]['approved_qty'] += payable_qty
        aggregation[code]['payable_usd'] += (payable_qty * price)

        if executed > payable_qty:
            aggregation[code]['blocked_usd'] += ((executed - payable_qty) * price)

    row = 6
    total_project_payable = 0.0
    total_project_blocked = 0.0

    for code in sorted(aggregation.keys()):
        data = aggregation[code]
        dict_item = dictionary.get(code)

        ws.cell(row=row, column=1, value=code).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=2, value=dict_item.description_ar if dict_item else '').alignment = Alignment(horizontal='right')
        ws.cell(row=row, column=3, value=dict_item.unit if dict_item else '').alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=4, value=dict_item.unit_price if dict_item else 0.0).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=5, value=data['planned']).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=6, value=data['executed']).alignment = Alignment(horizontal='center')

        app_cell = ws.cell(row=row, column=7, value=round(data['approved_qty'], 2))
        app_cell.alignment = Alignment(horizontal='center')
        app_cell.font = Font(color="006100" if data['approved_qty'] == data['planned'] else "9C5700")

        pay_cell = ws.cell(row=row, column=8, value=round(data['payable_usd'], 2))
        pay_cell.alignment = Alignment(horizontal='center')
        pay_cell.font = Font(bold=True)

        block_cell = ws.cell(row=row, column=9, value=round(data['blocked_usd'], 2))
        block_cell.alignment = Alignment(horizontal='center')
        if data['blocked_usd'] > 0:
            block_cell.font = Font(color="9C0006", bold=True)
            block_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        total_project_payable += data['payable_usd']
        total_project_blocked += data['blocked_usd']

        for col in range(1, 10):
            ws.cell(row=row, column=col).border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        row += 1

    # صف الإجمالي
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    tot_label = ws.cell(row=row, column=1, value="GRAND TOTAL (USD)")
    tot_label.font = Font(bold=True, size=12)
    tot_label.alignment = Alignment(horizontal='right')

    tot_pay = ws.cell(row=row, column=8, value=round(total_project_payable, 2))
    tot_pay.font = Font(bold=True, size=12, color="006100")
    tot_pay.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    tot_block = ws.cell(row=row, column=9, value=round(total_project_blocked, 2))
    tot_block.font = Font(bold=True, size=12, color="9C0006")
    if total_project_blocked > 0:
        tot_block.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    # عروض الأعمدة
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 18
    ws.column_dimensions['G'].width = 22
    ws.column_dimensions['H'].width = 20
    ws.column_dimensions['I'].width = 20

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_ipc_pdf(db: Session) -> bytes:
    """نسخة PDF مختصرة من تقرير IPC"""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1.5*cm, leftMargin=1.5*cm)
    elements = []
    styles = getSampleStyleSheet()

    elements.append(Paragraph("Interim Payment Certificate (IPC) - Summary", styles['Heading1']))
    elements.append(Paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d')}", styles['Normal']))
    elements.append(Spacer(1, 10))

    dictionary = {d.boq_code: d for d in db.query(models.BoqDictionary).filter(models.BoqDictionary.is_active == True).all()}
    items = db.query(models.BoqItem).all()
    decisions = {d.boq_item_id: d for d in db.query(models.DecisionRecord).all()}

    aggregation = {}
    for item in items:
        code = item.boq_code
        if code not in aggregation:
            aggregation[code] = {'planned': 0.0, 'executed': 0.0, 'approved_qty': 0.0, 'payable_usd': 0.0, 'blocked_usd': 0.0}
        dict_item = dictionary.get(code)
        price = dict_item.unit_price if dict_item else 0.0
        planned = item.planned_qty or 0.0
        executed = item.achieved_qty or 0.0
        aggregation[code]['planned'] += planned
        aggregation[code]['executed'] += executed
        decision = decisions.get(item.id)
        payment_pct = decision.human_payment_pct if decision and decision.human_payment_pct is not None else (decision.system_payment_pct if decision else 0.0)
        payable_qty = planned * (payment_pct / 100.0)
        aggregation[code]['approved_qty'] += payable_qty
        aggregation[code]['payable_usd'] += (payable_qty * price)
        if executed > payable_qty:
            aggregation[code]['blocked_usd'] += ((executed - payable_qty) * price)

    data = [['BoQ Code', 'Planned Qty', 'Executed Qty', 'Approved Qty', 'Payable ($)', 'Blocked ($)']]
    total_payable = 0.0
    total_blocked = 0.0
    for code in sorted(aggregation.keys()):
        d = aggregation[code]
        data.append([code, str(d['planned']), str(d['executed']), str(round(d['approved_qty'],2)), str(round(d['payable_usd'],2)), str(round(d['blocked_usd'],2))])
        total_payable += d['payable_usd']
        total_blocked += d['blocked_usd']
    data.append(['TOTAL', '', '', '', str(round(total_payable,2)), str(round(total_blocked,2))])

    table = Table(data, colWidths=[3*cm, 2.5*cm, 2.5*cm, 2.5*cm, 3*cm, 3*cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#f2f2f2')),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()

# ==========================================
# Remarks Report (PDF & Excel)
# ==========================================
def generate_remarks_pdf(db: Session, from_date: datetime = None, to_date: datetime = None) -> bytes:
    """سجل ملاحظات الجودة PDF - النسخة الكاملة"""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], alignment=1, fontSize=18, textColor=colors.HexColor('#1F4E78'), spaceAfter=20)
    section_style = ParagraphStyle('SectionTitle', parent=styles['Heading2'], fontSize=14, textColor=colors.HexColor('#2c3e50'), spaceAfter=10, spaceBefore=15)

    elements.append(Paragraph("Quality Control Remarks Log", title_style))
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 20))

    query = db.query(models.Remark).join(models.Latrine)
    if from_date:
        query = query.filter(models.Remark.date_logged >= from_date)
    if to_date:
        query = query.filter(models.Remark.date_logged <= to_date)
    open_remarks = query.filter(models.Remark.status == 'open').order_by(models.Remark.date_logged.desc()).all()

    if open_remarks:
        elements.append(Paragraph(f"OPEN REMARKS ({len(open_remarks)})", section_style))
        for remark in open_remarks:
            data = [
                ['Latrine ID', remark.latrine.latrine_id if remark.latrine else 'N/A'],
                ['BoQ Code', remark.boq_code or 'General'],
                ['Severity', remark.severity.upper()],
                ['Description', remark.description or ''],
            ]
            table = Table(data, colWidths=[4*cm, 12*cm])
            bg_color = colors.HexColor('#FFC7CE') if remark.severity == 'critical' else colors.HexColor('#FFEB9C') if remark.severity == 'major' else colors.HexColor('#e8f4f8')
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), bg_color),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ]))
            elements.append(table)
            elements.append(Spacer(1, 10))
    else:
        elements.append(Paragraph("No open remarks found.", styles['Normal']))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()

def generate_remarks_excel(db: Session, from_date: datetime = None, to_date: datetime = None) -> bytes:
    """سجل ملاحظات الجودة Excel"""
    buffer = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Remarks Log"
    ws.sheet_view.rightToLeft = True

    headers = ['Remark ID', 'Latrine ID', 'BoQ Code', 'Severity', 'Status', 'Description', 'Date Logged']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell.alignment = Alignment(horizontal='center')

    query = db.query(models.Remark).join(models.Latrine)
    if from_date:
        query = query.filter(models.Remark.date_logged >= from_date)
    if to_date:
        query = query.filter(models.Remark.date_logged <= to_date)
    remarks = query.order_by(models.Remark.date_logged.desc()).all()

    for i, remark in enumerate(remarks, start=2):
        ws.cell(row=i, column=1, value=remark.remark_id or f"REM-{remark.local_uuid[:8]}" if remark.local_uuid else str(remark.id))
        ws.cell(row=i, column=2, value=remark.latrine.latrine_id if remark.latrine else 'N/A')
        ws.cell(row=i, column=3, value=remark.boq_code or 'General')
        ws.cell(row=i, column=4, value=remark.severity)
        ws.cell(row=i, column=5, value=remark.status)
        ws.cell(row=i, column=6, value=remark.description or '')
        ws.cell(row=i, column=7, value=remark.date_logged.strftime('%Y-%m-%d') if remark.date_logged else '')

    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 50
    ws.column_dimensions['G'].width = 15

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# ==========================================
# Daily Logs Report (PDF & Excel)
# ==========================================
def generate_daily_logs_pdf(db: Session, from_date: datetime = None, to_date: datetime = None) -> bytes:
    """تقرير يوميات الموقع PDF - النسخة الكاملة"""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1.5*cm, leftMargin=1.5*cm)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], alignment=1, fontSize=18, textColor=colors.HexColor('#1F4E78'), spaceAfter=10)
    subtitle_style = ParagraphStyle('CustomSubtitle', parent=styles['Normal'], alignment=1, fontSize=11, textColor=colors.HexColor('#7f8c8d'), spaceAfter=20)

    elements.append(Paragraph("Site Diary & Daily Operations Log", title_style))
    date_str = "All Records"
    if from_date and to_date:
        date_str = f"Period: {from_date.strftime('%Y-%m-%d')} to {to_date.strftime('%Y-%m-%d')}"
    elements.append(Paragraph(date_str, subtitle_style))

    query = db.query(models.DailyLog).order_by(models.DailyLog.date.desc())
    if from_date: query = query.filter(models.DailyLog.date >= from_date)
    if to_date: query = query.filter(models.DailyLog.date <= to_date)
    logs = query.all()

    if not logs:
        elements.append(Paragraph("No daily logs found.", styles['Normal']))
        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    data = [['Date', 'Engineer', 'Weather', 'Manpower', 'Inspected', 'Accepted', 'Remarks', 'Equipment / Notes']]
    for log in logs:
        notes_text = ""
        if log.equipment: notes_text += f"Eq: {log.equipment}\n"
        if log.notes: notes_text += f"Note: {log.notes}"
        data.append([
            log.date.strftime('%Y-%m-%d'),
            log.engineer or '-',
            log.weather or '-',
            str(log.manpower or 0),
            str(log.latrines_inspected or 0),
            str(log.latrines_accepted or 0),
            str(log.remarks_issued or 0),
            notes_text or '-'
        ])

    col_widths = [2.5*cm, 3.5*cm, 2*cm, 2*cm, 2*cm, 2*cm, 2*cm, 10*cm]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (7, 1), (7, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()

def generate_daily_logs_excel(db: Session, from_date: datetime = None, to_date: datetime = None) -> bytes:
    """يوميات الموقع Excel"""
    buffer = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Site Diary"
    ws.sheet_view.rightToLeft = True

    headers = ['Date', 'Engineer', 'Weather', 'Manpower', 'Inspected', 'Accepted', 'Remarks', 'Equipment', 'Notes']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell.alignment = Alignment(horizontal='center')

    query = db.query(models.DailyLog).order_by(models.DailyLog.date.desc())
    if from_date: query = query.filter(models.DailyLog.date >= from_date)
    if to_date: query = query.filter(models.DailyLog.date <= to_date)
    logs = query.all()

    for i, log in enumerate(logs, start=2):
        ws.cell(row=i, column=1, value=log.date.strftime('%Y-%m-%d') if log.date else '')
        ws.cell(row=i, column=2, value=log.engineer or '')
        ws.cell(row=i, column=3, value=log.weather or '')
        ws.cell(row=i, column=4, value=log.manpower or 0)
        ws.cell(row=i, column=5, value=log.latrines_inspected or 0)
        ws.cell(row=i, column=6, value=log.latrines_accepted or 0)
        ws.cell(row=i, column=7, value=log.remarks_issued or 0)
        ws.cell(row=i, column=8, value=log.equipment or '')
        ws.cell(row=i, column=9, value=log.notes or '')

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 10
    ws.column_dimensions['H'].width = 30
    ws.column_dimensions['I'].width = 30

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# ==========================================
# Matrix Report (Excel & PDF)
# ==========================================
def generate_matrix_excel(db: Session) -> bytes:
    """تقرير أفقي يعرض كل مستفيد في صف، وكل بند كعمودين (كمية/سعر)"""
    buffer = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Progress Matrix"
    ws.sheet_view.rightToLeft = True

    dictionary = db.query(models.BoqDictionary).filter(models.BoqDictionary.is_active == True).order_by(models.BoqDictionary.boq_code).all()
    boq_codes = [d.boq_code for d in dictionary]

    ws.merge_cells('A1:D1')
    ws['A1'] = "Project Progress Matrix (Horizontal View)"
    ws['A1'].font = Font(size=14, bold=True)

    ws.cell(row=2, column=1, value="رقم الحمام").font = Font(bold=True)
    ws.cell(row=2, column=2, value="المستفيد (مشفر)").font = Font(bold=True)
    ws.cell(row=2, column=3, value="المربع").font = Font(bold=True)
    ws.cell(row=2, column=4, value="إجمالي الإنجاز %").font = Font(bold=True)

    col_idx = 5
    for d in dictionary:
        ws.merge_cells(start_row=1, start_column=col_idx, end_row=1, end_column=col_idx+1)
        header_cell = ws.cell(row=1, column=col_idx, value=f"{d.boq_code} - {d.description_ar}")
        header_cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_cell.font = Font(color="FFFFFF", bold=True)
        header_cell.alignment = Alignment(horizontal='center')
        ws.cell(row=2, column=col_idx, value="الكمية المنفذة").font = Font(bold=True)
        ws.cell(row=2, column=col_idx+1, value="التكلفة ($)").font = Font(bold=True)
        col_idx += 2

    latrines = db.query(models.Latrine).all()
    all_items = db.query(models.BoqItem).all()
    items_map = {}
    for item in all_items:
        if item.latrine_id not in items_map:
            items_map[item.latrine_id] = {}
        items_map[item.latrine_id][item.boq_code] = item

    row_idx = 3
    for latrine in latrines:
        ws.cell(row=row_idx, column=1, value=latrine.latrine_id)
        ws.cell(row=row_idx, column=2, value="[Encrypted]")
        ws.cell(row=row_idx, column=3, value=latrine.block_no)
        ws.cell(row=row_idx, column=4, value=f"{latrine.overall_pct}%")
        col_idx = 5
        for d in dictionary:
            item = items_map.get(latrine.id, {}).get(d.boq_code)
            qty = item.achieved_qty if item else 0.0
            cost = qty * d.unit_price
            ws.cell(row=row_idx, column=col_idx, value=qty)
            ws.cell(row=row_idx, column=col_idx+1, value=cost)
            col_idx += 2
        row_idx += 1

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_matrix_pdf(db: Session) -> bytes:
    """نسخة PDF مختصرة من مصفوفة الإنجاز (جدول ملخص)"""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1*cm, leftMargin=1*cm)
    elements = []
    styles = getSampleStyleSheet()

    elements.append(Paragraph("Project Progress Matrix - Summary", styles['Heading1']))
    elements.append(Spacer(1, 10))

    dictionary = db.query(models.BoqDictionary).filter(models.BoqDictionary.is_active == True).order_by(models.BoqDictionary.boq_code).all()
    latrines = db.query(models.Latrine).all()
    all_items = db.query(models.BoqItem).all()
    items_map = {}
    for item in all_items:
        if item.latrine_id not in items_map:
            items_map[item.latrine_id] = {}
        items_map[item.latrine_id][item.boq_code] = item

    headers = ['Latrine ID', 'Block'] + [f"{d.boq_code}" for d in dictionary] + ['Overall %']
    data = [headers]
    for latrine in latrines:
        row = [latrine.latrine_id, latrine.block_no]
        for d in dictionary:
            item = items_map.get(latrine.id, {}).get(d.boq_code)
            achieved = item.achieved_qty if item else 0
            row.append(str(achieved))
        row.append(f"{latrine.overall_pct}%")
        data.append(row)

    col_widths = [2.5*cm, 1.5*cm] + [1.2*cm]*len(dictionary) + [2*cm]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 7),
        ('GRID', (0,0), (-1,-1), 0.3, colors.grey),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()